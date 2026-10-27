"""
Zet alle ingevulde Excel-bestanden in de map excel/ om naar data/weekends.js
voor de website.

Gebruik (vanuit de projectmap):
    python tools/excel_naar_data.py

Vereist: pip install openpyxl
Let op: sla de Excel eerst op in Excel zelf, zodat de berekende waarden
(tabblad Afrekening) mee in het bestand bewaard zijn.
"""
import json
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
EXCEL_MAP = ROOT / "excel"
ASSETS = ROOT / "assets"
FOTO_TYPES = {".jpg", ".jpeg", ".png", ".webp"}
UITVOER = ROOT / "data" / "weekends.js"


def tekst(v):
    return str(v).strip() if v not in (None, "") else ""


def getal(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return 0.0


def datum(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, datetime) else None


def lees_weekend(pad):
    wb = load_workbook(pad, data_only=True)
    waarschuwingen = []

    # --- Info: veld in kolom A, waarde in kolom B ---
    info = {tekst(r[0]): r[1] for r in wb["Info"].iter_rows(min_row=2, values_only=True) if r[0]}
    link = tekst(info.get("Link huisje"))
    w = {
        "id": tekst(info.get("Weekend-ID")),
        "naam": tekst(info.get("Naam weekend")),
        "start": datum(info.get("Startdatum")),
        "eind": datum(info.get("Einddatum")),
        "locatie": tekst(info.get("Locatie")),
        "huisje": tekst(info.get("Naam huisje")),
        "linkHuisje": link if link.startswith("http") else "",
        "linkOpmerking": "" if link.startswith("http") else link,
        "opmerkingen": tekst(info.get("Opmerkingen")),
        "afgerond": tekst(info.get("Afrekening afgerond")) == "Ja",
        "totaal": getal(info.get("Totaal kosten")),
        "vooraf": getal(info.get("Waarvan vooraf betaald")),
    }

    # --- Afrekening: één rij per deelnemer tot de rij "Totaal" ---
    w["deelnemers"] = []
    for r in wb["Afrekening"].iter_rows(min_row=2, values_only=True):
        if tekst(r[0]) == "Totaal":
            break
        if not tekst(r[0]):
            continue
        w["deelnemers"].append({
            "naam": tekst(r[0]),
            "nachten": getal(r[1]),
            "vooraf": getal(r[2]),
            "voorgeschoten": getal(r[3]),
            "aandeel": getal(r[4]),
            "saldo": getal(r[5]),
            "voldaan": tekst(r[7]) == "Ja",
            "totaal": getal(r[8]),
        })

    # --- Kosten ---
    w["kosten"] = []
    for r in wb["Kosten"].iter_rows(min_row=2, values_only=True):
        if r[3] in (None, ""):
            continue
        w["kosten"].append({
            "datum": datum(r[0]),
            "omschrijving": tekst(r[1]),
            "categorie": tekst(r[2]) or "Andere",
            "bedrag": getal(r[3]),
            "betaaldDoor": tekst(r[4]),
            "verdeling": tekst(r[5]) or "Nachten",
            "vooraf": tekst(r[6]) == "Ja",
            "opmerking": tekst(r[7]),
        })

    # --- Foto's: assets/<weekend-id>/ ---
    # page_picture.* = verbleekte achtergrond van de pagina, de rest = sfeerbeelden
    w["achtergrond"] = ""
    w["fotos"] = []
    map_ = ASSETS / w["id"]
    if map_.is_dir():
        for f in sorted(map_.iterdir(), key=lambda p: p.name.lower()):
            if f.suffix.lower() not in FOTO_TYPES:
                continue
            # quote(): spaties en haakjes in bestandsnamen (bv. "WhatsApp Image ... (1).jpeg") veilig maken voor de browser
            pad_web = f"assets/{w['id']}/{quote(f.name)}"
            if f.stem.lower() == "page_picture":
                w["achtergrond"] = pad_web
            else:
                w["fotos"].append(pad_web)

    # --- Controles ---
    jaar = w["id"][:4]
    if not w["start"]:
        waarschuwingen.append("geen startdatum ingevuld")
    elif w["start"][:4] != jaar:
        waarschuwingen.append(f"startdatum {w['start']} past niet bij het jaar {jaar}")
    if w["totaal"] and not w["deelnemers"]:
        waarschuwingen.append("geen berekende waarden gevonden: open en bewaar het bestand eerst in Excel")
    if any(not k["categorie"] or k["categorie"] == "Andere" for k in w["kosten"]):
        waarschuwingen.append("kosten zonder categorie (worden getoond als 'Andere')")
    return w, waarschuwingen


def main():
    bestanden = sorted(p for p in EXCEL_MAP.glob("*.xlsx")
                       if "Template" not in p.name and not p.name.startswith("~$"))
    if not bestanden:
        sys.exit(f"Geen Excel-bestanden gevonden in {EXCEL_MAP}")

    weekends = []
    for pad in bestanden:
        w, waarschuwingen = lees_weekend(pad)
        weekends.append(w)
        status = "OK" if not waarschuwingen else "LET OP: " + "; ".join(waarschuwingen)
        print(f"{pad.name:36} {len(w['kosten']):3} kosten  {len(w['deelnemers'])} deeln.  "
              f"{len(w['fotos']):3} foto's  {'achtergrond' if w['achtergrond'] else 'geen achtergrond'}  {status}")

    UITVOER.parent.mkdir(exist_ok=True)
    UITVOER.write_text(
        "// AUTOMATISCH GEMAAKT door tools/excel_naar_data.py — niet met de hand aanpassen.\n"
        "// Aanpassen? Wijzig de Excel in excel/ en draai het script opnieuw.\n"
        "window.WEEKENDS = " + json.dumps(weekends, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    print(f"\n{len(weekends)} weekends weggeschreven naar {UITVOER.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
